"""Build one editable US-letter poster slide from poster.html.
Requires python-pptx and lxml. Run from any directory.
"""
from pathlib import Path
from lxml import html
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
DOC = html.fromstring((ROOT / 'poster.html').read_text())
PRS = Presentation()
PRS.slide_width, PRS.slide_height = Inches(8.5), Inches(11)
SLIDE = PRS.slides.add_slide(PRS.slide_layouts[6])
ACCENT, INK, PAPER, PALE = 'A52A2A', '222222', 'FFFFFF', 'F3F3F3'
SLIDE.background.fill.solid()
SLIDE.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)

def content(el):
    return ' '.join(el.text_content().split())

def select(cls):
    return DOC.xpath('//*[contains(concat(" ",normalize-space(@class)," ")," '+cls+' ")]')

def box(x,y,w,h,fill=None,line=None,rounded=False):
    s=SLIDE.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    if rounded:
        s.adjustments[0]=0.08
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb=RGBColor.from_string(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb=RGBColor.from_string(line); s.line.width=Pt(.6)
    else: s.line.fill.background()
    s._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in s._element.xpath('.//a:effectRef'):
        effect.set('idx', '0')
    return s

def text(value,x,y,w,h,size=7,color=INK,bold=False,align=None,font='Arial'):
    s=SLIDE.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
    tf=s.text_frame
    tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_before=p.space_after=Pt(0); p.line_spacing=1.08
        if align is not None: p.alignment=align
    return s

box(18,18,576,3,ACCENT)
text('CSC 510  /  PROJECT 1b  /  PROPOSAL',18,31,540,12,7,'666666')
text('Weeklies',18,49,530,40,34,INK,True,font='Georgia')
text('Your meals. Your schedule. Your budget.',18,94,530,22,16,INK)
text('Group 11   |   Srikar Desemsetti, Rohan Patel, Alex Tanton, Kushal Upretti',18,121,530,10,7,'555555')
box(18,138,576,.6,'CCCCCC')
mission=select('mission')[0]
text(content(mission.find('h2')).upper(),18,145,576,13,9,bold=True)
text(content(mission.find('p')),18,160,576,45,7.3)
for i,el in enumerate(select('stakeholders')[0]):
    x=18+i*146
    box(x,211,1,33,'CCCCCC')
    text(content(el.find('h3')),x+7,211,130,11,8,bold=True)
    text(content(el.find('p')),x+7,224,130,25,6.7)
for i,el in enumerate(select('sell')[0]):
    x=18+i*195
    box(x,255,186,1,ACCENT)
    text(content(el.find('div')),x+8,263,170,10,7,ACCENT,True)
    text(content(el.find('h3')),x+8,277,170,12,7.6,bold=True)
    text(content(el.find('p')),x+8,292,170,30,6.8)
text('DEVELOPMENT PLAN',18,337,340,12,9,bold=True)
text('Proposed scope · goals paired with methods',385,339,209,10,6,align=PP_ALIGN.RIGHT)
for i,el in enumerate(select('phase')):
    x=18+i*195
    box(x,355,186,23,PALE)
    box(x,355,186,1,ACCENT if i==1 else 'AAAAAA')
    text(content(el.find('h2')).upper(),x+9,364,165,12,10,bold=True)
    text(content(el.find('div')),x+9,379,165,10,5.8,'666666')
    for j,m in enumerate(el.xpath('./div[@class="milestone"]')):
        y=395+j*40
        text(content(m.find('h3')),x+9,y,168,10,7.5,bold=True)
        text(content(m.find('p')),x+9,y+12,168,28,6.35)
box(18,532,576,42,PALE)
for i,(n,label) in enumerate([('28','documented use cases'),('28','Project 1a test modules'),('172','test functions')]):
    x=23+i*74
    text(n,x,538,72,19,17,INK,True,PP_ALIGN.CENTER)
    text(label,x,559,72,8,5.5,INK,align=PP_ALIGN.CENTER)
text('Counts checked in use-cases.md and proj2/sef26tests/.\nSource inventory only. Execution, pass/fail totals, and coverage still need verification.',260,542,321,26,6,INK)
for i,el in enumerate(select('screens')[0]):
    x=18+i*195
    box(x,585,186,50,'FAFAFA','BBBBBB')
    ph=el.find('div')
    title=content(ph.find('b'))
    subtitle=ph.find('b').tail.strip()
    text(title,x+5,596,176,10,6.4,bold=True,align=PP_ALIGN.CENTER)
    text(subtitle,x+5,608,176,10,6.5,align=PP_ALIGN.CENTER)
    text(content(ph.find('small')),x+5,620,176,9,5.5,'666666',align=PP_ALIGN.CENTER)
    text(content(el.find('figcaption')),x,640,186,10,5.6)
for i,el in enumerate(select('tech')):
    x=18+i*117
    box(x,662,16,16,PALE)
    text(content(el.find('span')),x,665,16,10,6,INK,True,PP_ALIGN.CENTER)
    text(content(el.find('strong')),x+20,665,94,10,5.9,bold=True)
    text(content(el.find('p')),x,683,109,21,5.8)
for i,el in enumerate(select('links')[0]):
    x=18+i*195
    text(content(el.find('h3')),x,712,186,10,6.5,bold=True,align=PP_ALIGN.CENTER)
    box(x+74,726,38,38,None,'BBBBBB',False)
    text('INSERT\n'+['REPO QR','FORUM QR','VIDEO QR'][i],x+74,736,38,20,5.5,'666666',align=PP_ALIGN.CENTER)
    url=el.find('a') if el.find('a') is not None else el.find('span')
    s=text(content(url),x,768,186,9,5.5,align=PP_ALIGN.CENTER)
    if url.get('href'):
        s.text_frame.paragraphs[0].runs[0].hyperlink.address=url.get('href')
text('DRAFT · Replace team details, screenshots, test evidence, icon badges, and QR/link placeholders before submission.',18,782,576,7,5,'777777',align=PP_ALIGN.CENTER)
SLIDE.notes_slide.notes_text_frame.text=(ROOT/'poster.md').read_text()
PRS.core_properties.title='Weeklies — Editable poster draft'
PRS.core_properties.subject='Project 1b proposed next version; letter-size portrait poster'
PRS.save(ROOT/'weeklies-poster.pptx')
print(f'Saved {ROOT / "weeklies-poster.pptx"}: {len(SLIDE.shapes)} editable shapes/text boxes')
