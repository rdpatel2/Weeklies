from pptx import Presentation
from pptx.util import Inches

path = "poster-draft/poster.pptx"
prs = Presentation(path)
slide = prs.slides[0]

# The evidence band contains six metric value/label pairs. Align each pair
# to the center of one equal-width column in the band.
metrics = {
    "28": [],
    "172": [],
    "164": [],
    "8": [],
    "75.8%": [],
}
values = []
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    label = " ".join(shape.text.split())
    if label in {"28", "172", "164", "8", "75.8%"}:
        values.append(shape)

# Preserve the existing left-to-right order, including the duplicate 28.
values.sort(key=lambda s: s.left)
labels = []
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    label = " ".join(shape.text.split())
    if label in {"documented use cases", "Project 1a test modules", "test functions", "Tests passing", "Tests skipped (temporary)", "Test Coverage"}:
        labels.append(shape)
labels.sort(key=lambda s: s.left)

left = Inches(0.19)
column = Inches(0.88)
width = Inches(0.78)
for i, shape in enumerate(values):
    shape.left = left + i * column
    shape.width = width
for i, shape in enumerate(labels):
    shape.left = left + i * column
    shape.width = width

# Give the percentage enough room to stay on one line at poster scale.
for shape in values:
    if " ".join(shape.text.split()) == "75.8%":
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = run.font.size * 0.82

prs.save(path)
print(f"Aligned {len(values)} metric values and {len(labels)} metric labels in {path}")
